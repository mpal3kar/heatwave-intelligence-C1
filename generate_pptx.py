import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette
    NAVY = RGBColor(30, 41, 59)        # #1e293b
    CRIMSON = RGBColor(192, 57, 43)    # #c0392b
    AMBER = RGBColor(230, 126, 34)     # #e67e22
    LIGHT_BG = RGBColor(244, 246, 249) # #f4f6f9
    WHITE = RGBColor(255, 255, 255)
    DARK_GRAY = RGBColor(51, 65, 85)
    MUTED = RGBColor(100, 116, 139)
    BORDER_COL = RGBColor(203, 213, 225)

    def add_header(slide, title_text, category_text="MINI PROJECT PRESENTATION"):
        # Top banner background
        top_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        top_rect.fill.solid()
        top_rect.fill.fore_color.rgb = NAVY
        top_rect.line.fill.background()

        # Accent line under banner
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = AMBER
        accent.line.fill.background()

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.3))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = AMBER

        # Main Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.5), Inches(0.65))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = WHITE

        # Footer
        foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.35))
        tf_f = foot_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = "Mini Project Evaluation | Team: Mohammad Palekar, Gaurav, Samarth"
        p_f.font.size = Pt(10)
        p_f.font.color.rgb = MUTED

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_COL):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = LIGHT_BG
    bg1.line.fill.background()

    card1 = add_card(s1, 1.2, 0.9, 10.933, 5.7, WHITE, RGBColor(226, 232, 240))

    tb = s1.shapes.add_textbox(Inches(1.6), Inches(1.2), Inches(10.1), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "MINI PROJECT PRESENTATION"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = MUTED
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf.add_paragraph()
    p1.text = "Academic Year 2026-27 • Semester III"
    p1.font.size = Pt(12)
    p1.font.color.rgb = DARK_GRAY
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(20)

    p2 = tf.add_paragraph()
    p2.text = "Climate Intelligence for Heatwave Monitoring,\nPrediction & Early-Warning Web Portal"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = CRIMSON
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "System Architecture & Implementation"
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = NAVY
    p3.alignment = PP_ALIGN.CENTER
    p3.space_after = Pt(24)

    p4 = tf.add_paragraph()
    p4.text = "TEAM MEMBERS (DIV: C-1 | BATCH: C-1 | SEM: III | AY: 2026-27)"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = AMBER
    p4.alignment = PP_ALIGN.CENTER
    p4.space_after = Pt(8)

    p5 = tf.add_paragraph()
    p5.text = "1. Mohammad Palekar (16010125167)  •  2. Gaurav (16010125164)  •  3. Samarth (16010125161)"
    p5.font.size = Pt(15)
    p5.font.bold = True
    p5.font.color.rgb = NAVY
    p5.alignment = PP_ALIGN.CENTER
    p5.space_after = Pt(16)

    p6 = tf.add_paragraph()
    p6.text = "Evaluation Date: 8th October, Thursday (Lab Hours)"
    p6.font.size = Pt(13)
    p6.font.italic = True
    p6.font.color.rgb = MUTED
    p6.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 2: RUBRICS COMPLIANCE CHECKLIST
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Mini Project Rubrics & Evaluation Criteria Alignment")

    rubrics = [
        ("1. Timely Submission & Team Performance", "5 Marks",
         "• Submission on time before deadline (8th October).\n• Active participation of all 3 team members (Mohammad, Gaurav, Samarth).\n• Seamless integration of HTML5, CSS3, and JS modules into one unified web portal.",
         RGBColor(238, 242, 255), RGBColor(79, 70, 229)),
        ("2. Designing & Documentation", "10 Marks",
         "• Well-organized semantic content structure across 8 pages.\n• Proper GUI design properties: cohesive color scheme, CSS Box model, zebra tables, card layouts.\n• Comprehensive documentation & code compendium for each member.",
         RGBColor(254, 242, 242), RGBColor(220, 38, 38)),
        ("3. Accessibility & Layout Flexibility", "5 Marks",
         "• 100% Cross-browser compatible (tested on Chrome, Edge, Firefox).\n• Layout does not depend on fixed screen resolutions (responsive CSS flexbox & media queries).\n• Accessible tags (<abbr>, alt texts, <noscript> fallback).",
         RGBColor(240, 253, 244), RGBColor(22, 163, 74)),
        ("4. Presentation & Specific Features", "5 Marks",
         "• Live demonstration of specific features: JS Regex validation, Climate Array Analytics, Object methods, Image Maps, Multimedia.\n• Clear presentation answering all technical and viva questions.",
         RGBColor(254, 243, 199), RGBColor(217, 119, 6))
    ]

    for idx, (title, marks, body, bg_col, accent_col) in enumerate(rubrics):
        x = 0.8 + (idx % 2) * 5.95
        y = 1.45 + (idx // 2) * 2.7
        add_card(s2, x, y, 5.75, 2.5, bg_col, accent_col)
        tb_r = s2.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.15), Inches(5.45), Inches(2.2))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_t = tf_r.paragraphs[0]
        p_t.text = f"{title} [{marks}]"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = accent_col
        p_t.space_after = Pt(8)

        for line in body.split("\n"):
            p_b = tf_r.add_paragraph()
            p_b.text = line
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = DARK_GRAY
            p_b.space_after = Pt(3)

    # ==========================================
    # SLIDE 3: PROBLEM STATEMENT & BACKGROUND
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Problem Statement & Context")

    add_card(s3, 0.8, 1.45, 11.733, 5.35)
    tb3 = s3.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(11.1), Inches(4.8))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "Problem Statement: Climate Intelligence Heatwave Monitoring & Early-Warning Web Portal"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p.space_after = Pt(14)

    bullets = [
        "Rising Global & Regional Temperatures: Heatwave frequency and severity have surged drastically across urban and regional sectors, posing direct health risks including heatstroke and dehydration.",
        "Need for Multi-Stakeholder Intelligence: Meteorological agencies, disaster management response officers, and citizens lack a single lightweight, zero-latency dashboard that visualizes real-time heat indices, forecasts, and actionable bulletins.",
        "Disaster Management Gaps: Delays in early warnings and communication bottlenecks between Automated Weather Stations (AWS) and public advisory channels lead to preventable casualties.",
        "Target Solution: A fully functional, standards-compliant web portal front-end built using HTML5, styled with modular CSS3, and empowered with client-side JavaScript for threshold calculations and strict regex-based enrollment validation."
    ]

    for b in bullets:
        p_b = tf3.add_paragraph()
        parts = b.split(":", 1)
        p_b.text = "• " + parts[0] + ":"
        p_b.font.bold = True
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = NAVY
        
        run = p_b.add_run()
        run.text = parts[1]
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY
        p_b.space_after = Pt(12)

    # ==========================================
    # SLIDE 4: OBJECTIVES OF THE MINI PROJECT
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Objectives & Learning Outcomes")

    objs = [
        ("Semantic HTML5 Architecture", "Design structured, accessible multi-page documents employing proper semantic tags (<header>, <nav>, <main>, <section>, <article>, <aside>, <footer>) for climatological data."),
        ("Modern CSS3 Styling & Layouts", "Implement separation of structure and presentation using external stylesheets, CSS box model, responsive flexbox layouts, zebra-striped tables, and keyframe animation effects."),
        ("Robust Client-Side Form Validation", "Construct an error-preventing registration engine using JavaScript Regular Expressions for Indian phone numbers, emails, observer names, AWS IDs, and postal codes."),
        ("Interactive Climate Analytics Engine", "Perform mathematical temperature threshold difference calculations, multi-day array traversals (min/max/average), and object-oriented encapsulation using JavaScript objects and methods.")
    ]

    for i, (title, desc) in enumerate(objs):
        x = 0.8 + (i % 2) * 5.95
        y = 1.45 + (i // 2) * 2.7
        add_card(s4, x, y, 5.75, 2.5)
        tb_o = s4.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.2), Inches(5.35), Inches(2.1))
        tf_o = tb_o.text_frame
        tf_o.word_wrap = True

        p_t = tf_o.paragraphs[0]
        p_t.text = f"{i+1}. {title}"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = CRIMSON
        p_t.space_after = Pt(8)

        p_d = tf_o.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 5: METHODOLOGY & TECHNOLOGIES USED
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Methodology & Technologies Used")

    add_card(s5, 0.8, 1.45, 11.733, 5.35)
    tb5 = s5.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(11.1), Inches(4.8))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    tech_items = [
        ("HTML5 (HyperText Markup Language 5)", "Core skeleton providing semantic tags (<header>, <nav>, <article>, <aside>), text formatting (<mark>, <abbr>, <sup>, <sub>), client-side coordinate image maps (<map>, <area>), multi-row/col tables, and native media (<video>, <audio>, <iframe>, <noscript>)."),
        ("CSS3 (Cascading Style Sheets 3)", "Modular styling using external style.css, responsive flexbox layout, CSS Box Model (margin, border, padding, border-radius, box-shadow), pseudo-classes (:hover, :focus, :nth-child), and glowing CSS @keyframes animations."),
        ("Vanilla JavaScript (ES6+)", "Client-side computing without external libraries: DOM event listeners (DOMContentLoaded, submit, click), Regex pattern testing (test()), dynamic error injection via innerHTML, array traversal loops, and climate object literals with methods and this binding."),
        ("Separation of Concerns & Modularity", "Strict separation between Content (HTML files), Presentation (css/style.css), and Behavior (js/validation.js & js/analytics.js) ensuring high maintainability and W3C validation compliance.")
    ]

    for title, desc in tech_items:
        p_m = tf5.add_paragraph()
        p_m.text = "• " + title + ": "
        p_m.font.bold = True
        p_m.font.size = Pt(13)
        p_m.font.color.rgb = NAVY
        
        run = p_m.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY
        p_m.space_after = Pt(11)

    # ==========================================
    # SLIDE 6: TEAM RESPONSIBILITY & WORK DIVISION
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Team Structure & Work Distribution (3 Parts)")

    members = [
        ("Mohammad Palekar (16010125167)", "Project Lead & Core HTML5 Architecture",
         "• Developed index.html, about.html, and media.html.\n• Designed multi-column meteorological tables (forecast.html) with rowspan and colspan.\n• Implemented client-side Image Maps (hotspot.html) with clickable sector coordinates.\n• Embedded native HTML5 multimedia (<video>, <audio>, <iframe> radar feeds) and <noscript> fallbacks.",
         RGBColor(254, 242, 242), CRIMSON),
        ("Gaurav (16010125164)", "Front-End UI/UX & CSS3 Styling Specialist",
         "• Engineered global styling sheet (css/style.css) adhering to separation of concerns.\n• Built responsive layout grid with flexbox, card containers, and sticky navigation bar.\n• Styled data tables with zebra-striping (tr:nth-child(even)) and hover feedback.\n• Crafted CSS @keyframes glowing alert box for Level 4 extreme heat advisory notices.",
         RGBColor(240, 253, 244), RGBColor(22, 163, 74)),
        ("Samarth (16010125161)", "JavaScript Engine & Form Validation Engineer",
         "• Programmed js/validation.js utilizing Regular Expressions for 6 strict field constraints.\n• Created DOM error injection and real-time success alert banners on form submission.\n• Engineered js/analytics.js for multi-day temperature array analysis (min, max, average, threshold count).\n• Implemented object-oriented climateData entity with member methods using this keyword.",
         RGBColor(238, 242, 255), RGBColor(79, 70, 229))
    ]

    for i, (name, role, tasks, bg_col, accent_col) in enumerate(members):
        x = 0.8 + i * 3.95
        y = 1.45
        add_card(s6, x, y, 3.8, 5.35, bg_col, accent_col)
        tb_m = s6.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.2), Inches(3.5), Inches(4.9))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True

        p_n = tf_m.paragraphs[0]
        p_n.text = name
        p_n.font.size = Pt(15)
        p_n.font.bold = True
        p_n.font.color.rgb = accent_col

        p_r = tf_m.add_paragraph()
        p_r.text = role
        p_r.font.size = Pt(11)
        p_r.font.bold = True
        p_r.font.color.rgb = DARK_GRAY
        p_r.space_after = Pt(12)

        for line in tasks.split("\n"):
            p_t = tf_m.add_paragraph()
            p_t.text = line
            p_t.font.size = Pt(11)
            p_t.font.color.rgb = DARK_GRAY
            p_t.space_after = Pt(6)

    # ==========================================
    # SLIDE 7: IMPLEMENTATION & DEMO ARCHITECTURE
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "System Implementation & Module Flow")

    add_card(s7, 0.8, 1.45, 11.733, 5.35)
    tb7 = s7.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(11.1), Inches(4.8))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "Modular Component Flow & System Integration:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p.space_after = Pt(10)

    modules = [
        ("Layer 1: Surveillance & Landing (index.html, alerts.html)", "Presents active heat advisories, key sensor statistics (peak temperatures, monitored zones), and text-level formatting (<mark>, <abbr>, <sup>, <sub>)."),
        ("Layer 2: Synoptic Data & Mapping (forecast.html, hotspot.html)", "Visualizes multi-station temperature matrices using advanced table tags with row/col spanning, alongside clickable regional radar maps navigating to localized feeds."),
        ("Layer 3: Analytical Computational Engine (analysis.html)", "Processes raw temperature readings array [34..42] through JavaScript algorithms, displaying min/max/average, while evaluating OOP climate objects and threshold calculators."),
        ("Layer 4: Enrollment & Public Notification (register.html)", "Gathers citizen and observer registrations, strictly preventing malformed inputs through client-side RegEx checks before triggering confirmation banners."),
        ("Layer 5: Public Media & Awareness (media.html)", "Delivers emergency audio broadcasts and instruction videos via native HTML5 media tags with iframe radar embeds.")
    ]

    for title, desc in modules:
        p_mod = tf7.add_paragraph()
        p_mod.text = "• " + title + ": "
        p_mod.font.bold = True
        p_mod.font.size = Pt(12)
        p_mod.font.color.rgb = NAVY

        run = p_mod.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY
        p_mod.space_after = Pt(8)

    # ==========================================
    # SLIDE 8: RESULTS - HOME & ALERT BULLETIN (SCREENSHOTS)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Results / Outcome: Home Portal & Alert Bulletin")

    add_card(s8, 0.8, 1.45, 5.75, 4.2)
    s8.shapes.add_picture(r"d:\Programming\Heatwave-MiniProject\screenshots\index.png", Inches(0.9), Inches(1.55), Inches(5.55), Inches(3.95))
    tb8_l = s8.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(5.75), Inches(1.1))
    tf8_l = tb8_l.text_frame
    tf8_l.word_wrap = True
    p = tf8_l.paragraphs[0]
    p.text = "Figure 1: Home Dashboard (index.html)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p1 = tf8_l.add_paragraph()
    p1.text = "Features global attributes, glowing CSS alert box, quick status sidebar, and key sensor metrics."
    p1.font.size = Pt(11)
    p1.font.color.rgb = DARK_GRAY

    add_card(s8, 6.78, 1.45, 5.75, 4.2)
    s8.shapes.add_picture(r"d:\Programming\Heatwave-MiniProject\screenshots\alerts.png", Inches(6.88), Inches(1.55), Inches(5.55), Inches(3.95))
    tb8_r = s8.shapes.add_textbox(Inches(6.78), Inches(5.8), Inches(5.75), Inches(1.1))
    tf8_r = tb8_r.text_frame
    tf8_r.word_wrap = True
    p = tf8_r.paragraphs[0]
    p.text = "Figure 2: Severe Alert Bulletin (alerts.html)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p1 = tf8_r.add_paragraph()
    p1.text = "Implements text-level formatting (<mark>, <em>, <sup>, <sub>, <abbr>) and agency verification."
    p1.font.size = Pt(11)
    p1.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 9: RESULTS - FORECAST & JS ANALYTICS (SCREENSHOTS)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Results / Outcome: Forecast Matrix & Climate Analytics")

    add_card(s9, 0.8, 1.45, 5.75, 4.2)
    s9.shapes.add_picture(r"d:\Programming\Heatwave-MiniProject\screenshots\forecast.png", Inches(0.9), Inches(1.55), Inches(5.55), Inches(3.95))
    tb9_l = s9.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(5.75), Inches(1.1))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    p = tf9_l.paragraphs[0]
    p.text = "Figure 3: Forecast Summary Matrix (forecast.html)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p1 = tf9_l.add_paragraph()
    p1.text = "Structured table with <caption>, <thead>, rowspan, colspan, and CSS zebra-striped rows."
    p1.font.size = Pt(11)
    p1.font.color.rgb = DARK_GRAY

    add_card(s9, 6.78, 1.45, 5.75, 4.2)
    s9.shapes.add_picture(r"d:\Programming\Heatwave-MiniProject\screenshots\analysis.png", Inches(6.88), Inches(1.55), Inches(5.55), Inches(3.95))
    tb9_r = s9.shapes.add_textbox(Inches(6.78), Inches(5.8), Inches(5.75), Inches(1.1))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    p = tf9_r.paragraphs[0]
    p.text = "Figure 4: JavaScript Analytics Engine (analysis.html)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p1 = tf9_r.add_paragraph()
    p1.text = "Live client-side array processing (min/max/avg), threshold calculator, and OOP climate object output."
    p1.font.size = Pt(11)
    p1.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 10: RESULTS - REGISTRATION FORM & HOTSPOT MAP (SCREENSHOTS)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Results / Outcome: Form Validation & Hotspot Map")

    add_card(s10, 0.8, 1.45, 5.75, 4.2)
    s10.shapes.add_picture(r"d:\Programming\Heatwave-MiniProject\screenshots\register.png", Inches(0.9), Inches(1.55), Inches(5.55), Inches(3.95))
    tb10_l = s10.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(5.75), Inches(1.1))
    tf10_l = tb10_l.text_frame
    tf10_l.word_wrap = True
    p = tf10_l.paragraphs[0]
    p.text = "Figure 5: Form Validation Module (register.html)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p1 = tf10_l.add_paragraph()
    p1.text = "Client-side validation with Regex patterns, field-level error feedback, and success message."
    p1.font.size = Pt(11)
    p1.font.color.rgb = DARK_GRAY

    add_card(s10, 6.78, 1.45, 5.75, 4.2)
    s10.shapes.add_picture(r"d:\Programming\Heatwave-MiniProject\screenshots\hotspot.png", Inches(6.88), Inches(1.55), Inches(5.55), Inches(3.95))
    tb10_r = s10.shapes.add_textbox(Inches(6.78), Inches(5.8), Inches(5.75), Inches(1.1))
    tf10_r = tb10_r.text_frame
    tf10_r.word_wrap = True
    p = tf10_r.paragraphs[0]
    p.text = "Figure 6: Interactive Hotspot Map (hotspot.html)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p1 = tf10_r.add_paragraph()
    p1.text = "Client-side clickable image map (<map> & <area>) linking geographical coordinates to portal feeds."
    p1.font.size = Pt(11)
    p1.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 11: TECHNICAL HIGHLIGHTS & CONCEPTS TESTED
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Technical Concepts Tested Across Experiments 1 to 5")

    add_card(s11, 0.8, 1.45, 11.733, 5.35)
    tb11 = s11.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(11.1), Inches(4.8))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    experiments = [
        ("Experiment 1 (HTML5 Core)", "Document structure, semantic tags, global attributes (id, class, title, tooltip), block vs inline formatting, lists (ol, ul, nested, dl), hyperlinks (mailto, tel), image maps, table tags, form elements, media (<video>, <audio>, <iframe>, <noscript>)."),
        ("Experiment 2 (CSS3 Layout & Styling)", "Separation of concerns (style.css), CSS box model, responsive flexbox layout, table zebra-striping with :nth-child(even), form UI states (:focus, :hover), image transitions, and keyframe animations (@keyframes alertGlow)."),
        ("Experiment 4 (JavaScript Form Validation)", "DOM event handling (submit, preventDefault()), field-level error messages in <span> elements, and Regular Expressions for observer name, mobile number, email, and AWS station ID."),
        ("Experiment 5 (JS Climate Logic & Arrays)", "Variables, relational and logical operators, multi-branch conditional statements, JavaScript Object literals with member functions (this keyword), array traversal loops for min/max/average, and future-date validation.")
    ]

    for title, desc in experiments:
        p_e = tf11.add_paragraph()
        p_e.text = "• " + title + ": "
        p_e.font.bold = True
        p_e.font.size = Pt(13)
        p_e.font.color.rgb = CRIMSON

        run = p_e.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY
        p_e.space_after = Pt(11)

    # ==========================================
    # SLIDE 12: CONCLUSION & FUTURE SCOPE
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Conclusion & Future Scope")

    add_card(s12, 0.8, 1.45, 5.75, 5.35)
    tb12_l = s12.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.35), Inches(4.8))
    tf12_l = tb12_l.text_frame
    tf12_l.word_wrap = True

    p = tf12_l.paragraphs[0]
    p.text = "Project Conclusion"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p.space_after = Pt(10)

    concl_points = [
        "Successful Integration: Integrated all curriculum concepts across Experiments 1 to 5 into an authentic, production-ready climate intelligence portal.",
        "Zero Framework Overhead: Built entirely using pure HTML5, CSS3, and Vanilla JavaScript, ensuring instantaneous load times (<0.2s) and 100% browser compatibility.",
        "User-Centric & Accessible: Clean, intuitive UI designed for ease of use by meteorologists, disaster officers, and common citizens during severe heatwave emergencies.",
        "Team Cohesion: Collaborative execution with clear separation of duties between Mohammad, Gaurav, and Samarth."
    ]
    for pt in concl_points:
        p_c = tf12_l.add_paragraph()
        p_c.text = "• " + pt
        p_c.font.size = Pt(12)
        p_c.font.color.rgb = DARK_GRAY
        p_c.space_after = Pt(8)

    add_card(s12, 6.78, 1.45, 5.75, 5.35)
    tb12_r = s12.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.35), Inches(4.8))
    tf12_r = tb12_r.text_frame
    tf12_r.word_wrap = True

    p = tf12_r.paragraphs[0]
    p.text = "Future Scope & Enhancements"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_after = Pt(10)

    future_points = [
        "Real-Time IoT API Integration: Connect to live RESTful APIs and WebSocket feeds from municipal Automatic Weather Stations (AWS).",
        "Geospatial GIS Mapping: Upgrade static image maps to dynamic Leaflet.js or OpenLayers interactive heatmaps with GIS polygon layers.",
        "Push Notification Gateway: Implement Web Push API and SMS Twilio integration for automated real-time danger broadcasts.",
        "Progressive Web App (PWA): Add Service Workers and offline caching so citizens can access emergency manuals even without internet."
    ]
    for pt in future_points:
        p_f = tf12_r.add_paragraph()
        p_f.text = "• " + pt
        p_f.font.size = Pt(12)
        p_f.font.color.rgb = DARK_GRAY
        p_f.space_after = Pt(8)

    # ==========================================
    # SLIDE 13: THANK YOU & Q&A
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    bg13 = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg13.fill.solid()
    bg13.fill.fore_color.rgb = NAVY
    bg13.line.fill.background()

    card13 = add_card(s13, 2.0, 1.2, 9.333, 5.1, WHITE, RGBColor(226, 232, 240))
    tb13 = s13.shapes.add_textbox(Inches(2.3), Inches(1.5), Inches(8.733), Inches(4.5))
    tf13 = tb13.text_frame
    tf13.word_wrap = True

    p = tf13.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = CRIMSON
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(10)

    p1 = tf13.add_paragraph()
    p1.text = "Questions & Answers (Viva Defense Ready)"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(20)

    p2 = tf13.add_paragraph()
    p2.text = "Semester III Mini Project Evaluation"
    p2.font.size = Pt(14)
    p2.font.color.rgb = DARK_GRAY
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(20)

    p3 = tf13.add_paragraph()
    p3.text = "Team: Mohammad Palekar (16010125167)  |  Gaurav (16010125164)  |  Samarth (16010125161)"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = AMBER
    p3.alignment = PP_ALIGN.CENTER

    out_path = r"d:\Programming\Heatwave-MiniProject\Heatwave_MiniProject_Presentation.pptx"
    prs.save(out_path)
    print(f"Final Presentation saved successfully to {out_path}")

if __name__ == "__main__":
    create_deck()
